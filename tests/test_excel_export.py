"""Verify portable Excel evidence with an independent XLSX reader.

Use saved report fixtures to check accounting parity, formula caches, sheet
structure, and literal content. Mutated fixtures exercise unavailable telemetry
and FX without provider calls. These checks do not require a desktop runtime.
AI attribution: Generated with AI assistance by Northstar.
Copyright (c) 2026 Martin.Bechard@DevConsult.ca
"""

import json
import os
import subprocess
import sys
from decimal import ROUND_HALF_UP, Decimal
from pathlib import Path

import openpyxl
import pytest

from reporting.excel_data import workbook_data
from reporting.export_excel import export_workbook
from reporting.pricing import load_prices
from reporting.schema import Run

ROOT = Path(__file__).resolve().parents[1]


def saved_data(sample="tool_chat"):
    """Load recorded evidence without rerunning or altering checked-in reports."""
    directory = ROOT / "reports" / sample
    run = Run.model_validate_json((directory / "run.json").read_text())
    return workbook_data(run, load_prices(directory / "prices.json"))


@pytest.mark.parametrize(
    "sample", ["tool_chat", "subagent_chat", "thinking_agent", "context_budget"]
)
def test_excel_accounting_and_structure(tmp_path, sample):
    """Cached totals match Python and projections retain editable formulas."""
    data = saved_data(sample)
    path = export_workbook(data, tmp_path / "report.xlsx")
    values = openpyxl.load_workbook(path, data_only=True)
    formulas = openpyxl.load_workbook(path)
    assert values.sheetnames == ["Turns", "Execution tree", "Reference data"]
    expected = Decimal(str(data["expected_usd"]))
    assert values["Turns"]["F2"].value == pytest.approx(float(expected))
    fx = Decimal(str(data["prices"]["exchange"]["rate"]))
    assert values["Turns"]["G2"].value == float(
        (expected * fx * 100000).quantize(Decimal(1), rounding=ROUND_HALF_UP)
    )
    assert "ROUND(" in formulas["Turns"]["G2"].value
    assert "'Reference data'!$B$2" in formulas["Turns"]["G2"].value
    assert values["Reference data"]["B2"].value == 100000
    assert all(sheet.freeze_panes == "B2" for sheet in formulas)
    # Every model has one charge total, and turn/run formulas point at those
    # totals instead of summing duplicated context and execution-tree values.
    labels = [row[2].value for row in values["Turns"].iter_rows(min_row=2)]
    assert labels.count("Call total") == sum(
        e["step"]["kind"] == "model" for e in data["events"]
    )
    for row in formulas["Turns"].iter_rows(min_row=2):
        if row[2].value == "Call total":
            assert "COUNT(" in row[5].value
    assert formulas.calculation.fullCalcOnLoad
    values.close()
    formulas.close()


@pytest.mark.parametrize("missing", ["fx", "usage", "tariff"])
def test_excel_unknown_values_are_not_zero(tmp_path, missing):
    """Absent FX affects EUR only; absent usage/pricing prevents complete totals."""
    data = saved_data()
    if missing == "fx":
        data["prices"]["exchange"] = None
    elif missing == "tariff":
        data["prices"]["models"] = {}
    else:
        event = next(e for e in data["events"] if e["step"]["kind"] == "model")
        event["step"]["usage"] = None
        for cell in event["cells"]:
            cell["partial"] = True
    path = export_workbook(data, tmp_path / "unknown.xlsx")
    workbook = openpyxl.load_workbook(path, data_only=True)
    assert workbook["Turns"]["E2"].value == "Unknown"
    assert workbook["Turns"]["G2"].value == "Unknown"
    if missing == "fx":
        assert workbook["Turns"]["F2"].value == pytest.approx(
            float(data["expected_usd"])
        )
    else:
        assert workbook["Turns"]["F2"].value == "Unknown"
        assert workbook["Turns"]["H2"].value == "Unknown"
    workbook.close()


def test_excel_cli_needs_no_node_or_codex(tmp_path):
    """The installed Python entry point exports safely with no executable search path."""
    data = saved_data()
    event = next(e for e in data["events"] if e["step"]["kind"] == "model")
    payload = '=HYPERLINK("https://example.invalid", "untrusted content")'
    event["user_prompt"] = payload
    source = tmp_path / "data.json"
    source.write_text(json.dumps(data, default=str))
    target = tmp_path / "report.xlsx"
    result = subprocess.run(
        [sys.executable, "-m", "reporting.export_excel", str(source), str(target)],
        cwd=tmp_path,
        env={
            **os.environ,
            "PATH": "",
            "LG_EXCEL_RUNTIME": str(tmp_path / "absent"),
        },
        capture_output=True,
        text=True,
        check=False,
    )
    assert result.returncode == 0, result.stderr
    workbook = openpyxl.load_workbook(target)
    cells = [cell for row in workbook["Turns"] for cell in row if cell.value == payload]
    assert cells and all(cell.data_type == "s" for cell in cells)
    workbook.close()


def test_excel_formulas_recalculate_after_edits(tmp_path):
    """A separate Excel formula engine proves caches are not hiding bad formulas.

    Changing executions must scale before rounding. Changing FX must affect EUR
    but leave USD unchanged, including execution-tree subtotals.
    """
    import formulas

    data = saved_data()
    path = export_workbook(data, tmp_path / "editable.xlsx")
    model = formulas.ExcelModel().loads(str(path)).finish()
    prefix = "'[editable.xlsx]"
    # Recompute every generated formula independently before testing edits.
    # This covers category prices, turn totals, and tree subtotals, not just
    # the run's cached headline cost.
    baseline = model.calculate()
    cached = openpyxl.load_workbook(path, data_only=True)
    expressions = openpyxl.load_workbook(path)
    for sheet in expressions:
        for row in sheet:
            for cell in row:
                if cell.data_type != "f":
                    continue
                key = f"'[editable.xlsx]{sheet.title.upper()}'!{cell.coordinate}"
                computed = baseline[key].value[0, 0]
                expected_value = cached[sheet.title][cell.coordinate].value
                if isinstance(expected_value, (int, float)):
                    assert computed == pytest.approx(expected_value), key
                else:
                    assert computed == expected_value, key
    cached.close()
    expressions.close()
    outputs = [prefix + "TURNS'!F2", prefix + "TURNS'!G2", prefix + "TURNS'!H2"]
    actual = model.calculate(
        inputs={
            prefix + "REFERENCE DATA'!B2": 200000,
            prefix + "REFERENCE DATA'!B3": 0.5,
        },
        outputs=outputs,
    )
    expected = Decimal(str(data["expected_usd"]))
    assert actual[outputs[0]].value[0, 0] == pytest.approx(float(expected))
    assert actual[outputs[1]].value[0, 0] == float(
        (expected * Decimal("0.5") * 200000).quantize(
            Decimal(1), rounding=ROUND_HALF_UP
        )
    )
    assert actual[outputs[2]].value[0, 0] == float(
        (expected * 200000).quantize(Decimal(1), rounding=ROUND_HALF_UP)
    )


def test_excel_keeps_distinct_cache_write_lifetimes(tmp_path):
    """Anthropic's 5-minute and 1-hour writes retain different billed rates.

    A single editable reference rate cannot represent this mixture. The workbook
    must retain the precise Python charge instead of pricing both at five minutes.
    """
    from reporting.schema import Usage

    directory = ROOT / "reports/tool_chat"
    run = Run.model_validate_json((directory / "run.json").read_text())
    model = next(step for step in run.steps if step.kind == "model")
    model.provider = "anthropic"
    model.model = "claude-sonnet-4-6"
    model.usage = Usage(
        input_tokens=1000,
        output_tokens=100,
        cache_write=500,
        cache_write_5m=200,
        cache_write_1h=300,
    )
    prices = load_prices(directory / "prices.json")
    prices.models.update(
        load_prices(ROOT / "tests/fixtures/accounting_prices.json").models
    )
    data = workbook_data(run, prices)
    output = export_workbook(data, tmp_path / "cache-lifetimes.xlsx")
    workbook = openpyxl.load_workbook(output, data_only=True)
    charges = [
        row[5].value
        for row in workbook["Turns"].iter_rows(min_row=2)
        if row[0].value == "R1"
        and row[2].value == "Cache write"
        and row[5].value is not None
    ]
    assert charges == [pytest.approx((200 * 3.75 + 300 * 6) / 1_000_000)]
    assert workbook["Turns"]["F2"].value == pytest.approx(float(data["expected_usd"]))
    workbook.close()
